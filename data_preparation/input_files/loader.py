from __future__ import annotations

import csv
import io
import json
import logging
import re
import tempfile
import zipfile
from pathlib import Path

logger = logging.getLogger(__name__)

THAI_RE = re.compile(r"[\u0E00-\u0E7F]")
MOJIBAKE_RE = re.compile(r"(เธ|เน|โ€”|โ|๏|ใ€|¢|Å|Ò|ä|§|�)")


def detect_encoding(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        raw.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        return "cp874"


def decode_best(raw: bytes, encodings: list[str]) -> tuple[str, str]:
    detected = detect_encoding(raw)
    candidates = [detected, *encodings, "utf-8-sig", "utf-8", "cp874", "tis-620", "windows-874"]
    best_text = ""
    best_encoding = detected
    best_score = -10**9
    for encoding in dict.fromkeys(candidates):
        try:
            text = raw.decode(encoding, errors="replace")
        except LookupError:
            continue
        score = len(THAI_RE.findall(text)) - (text.count("�") * 20) - (len(MOJIBAKE_RE.findall(text)) * 8)
        if score > best_score:
            best_text = text
            best_encoding = encoding
            best_score = score
    return best_text, best_encoding


class CSVParser:
    def __init__(self, encodings: list[str]):
        self.encodings = encodings

    def parse(self, path: Path) -> list[dict]:
        text, encoding = decode_best(path.read_bytes(), self.encodings)
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        rows = [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]
        logger.info("CSV decoded with %s: %s rows", encoding, len(rows))
        return [{"_source_file": path.name, "_source_type": "csv", **row} for row in rows]


class XLSXParser:
    def parse(self, path: Path) -> list[dict]:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        records: list[dict] = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(value).strip() if value is not None else f"col_{i}" for i, value in enumerate(rows[0])]
            for row in rows[1:]:
                records.append(
                    {
                        "_source_file": path.name,
                        "_source_type": "xlsx",
                        "_sheet": sheet_name,
                        **{header: value for header, value in zip(headers, row)},
                    }
                )
        workbook.close()
        return records


class PDFParser:
    def __init__(self, max_pages: int = 500):
        self.max_pages = max_pages

    def parse(self, path: Path) -> list[dict]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        records: list[dict] = []
        for index, page in enumerate(reader.pages[: self.max_pages]):
            text = (page.extract_text() or "").strip()
            if text:
                records.append(
                    {
                        "_source_file": path.name,
                        "_source_type": "pdf",
                        "_page": index + 1,
                        "text": text,
                    }
                )
        return records


class JSONParser:
    def __init__(self, encodings: list[str]):
        self.encodings = encodings

    def parse(self, path: Path) -> list[dict]:
        text, _ = decode_best(path.read_bytes(), self.encodings)
        if path.suffix.lower() == ".jsonl":
            items = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            data = json.loads(text)
            if isinstance(data, list):
                items = data
            elif isinstance(data, dict):
                items = next((value for value in data.values() if isinstance(value, list)), [data])
            else:
                items = []
        return [
            {"_source_file": path.name, "_source_type": path.suffix.lower().strip("."), **item}
            for item in items
            if isinstance(item, dict)
        ]


class TXTParser:
    def __init__(self, encodings: list[str]):
        self.encodings = encodings

    def parse(self, path: Path) -> list[dict]:
        text, encoding = decode_best(path.read_bytes(), self.encodings)
        text = text.strip()
        logger.info("Text decoded with %s: %s chars", encoding, len(text))
        if not text:
            return []
        return [{"_source_file": path.name, "_source_type": path.suffix.lower().strip("."), "text": text}]


def extract_zip_recursive(path: Path, dest: Path, depth: int = 0, max_depth: int = 3) -> list[Path]:
    if depth > max_depth:
        logger.warning("ZIP depth exceeds %s: %s", max_depth, path.name)
        return []

    extracted: list[Path] = []
    try:
        with zipfile.ZipFile(path, "r") as archive:
            archive.extractall(dest)
            for name in archive.namelist():
                member = dest / name
                if member.suffix.lower() == ".zip" and member.is_file():
                    nested_dest = member.parent / member.stem
                    nested_dest.mkdir(parents=True, exist_ok=True)
                    extracted.extend(extract_zip_recursive(member, nested_dest, depth + 1, max_depth))
                else:
                    extracted.append(member)
    except zipfile.BadZipFile as exc:
        logger.error("Bad ZIP file %s: %s", path.name, exc)
    return extracted


class DataLoader:
    SUPPORTED = {".csv", ".xlsx", ".xls", ".pdf", ".json", ".jsonl", ".txt", ".md", ".zip"}

    def __init__(self, cfg: dict):
        input_cfg = cfg.get("input_files", {})
        self.encodings = input_cfg.get("encodings", ["utf-8", "utf-8-sig", "cp874", "tis-620", "windows-874"])
        self.max_pages = int(input_cfg.get("pdf_max_pages", 500))
        self.max_depth = int(input_cfg.get("zip_max_depth", 3))
        self._parsers = {
            ".csv": CSVParser(self.encodings),
            ".xlsx": XLSXParser(),
            ".xls": XLSXParser(),
            ".pdf": PDFParser(self.max_pages),
            ".json": JSONParser(self.encodings),
            ".jsonl": JSONParser(self.encodings),
            ".txt": TXTParser(self.encodings),
            ".md": TXTParser(self.encodings),
        }

    def load(self, path: Path) -> list[dict]:
        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED:
            logger.warning("Unsupported file type %s: %s", suffix, path.name)
            return []
        if suffix == ".zip":
            return self._load_zip(path)

        parser = self._parsers.get(suffix)
        if parser is None:
            return []
        logger.info("Parsing %s (%s)", path.name, suffix)
        return parser.parse(path)

    def _load_zip(self, path: Path) -> list[dict]:
        with tempfile.TemporaryDirectory() as temp_dir:
            files = extract_zip_recursive(path, Path(temp_dir), max_depth=self.max_depth)
            records: list[dict] = []
            for file_path in files:
                if not file_path.is_file() or file_path.suffix.lower() not in self.SUPPORTED - {".zip"}:
                    continue
                for record in self.load(file_path):
                    record["_zip_source"] = path.name
                    records.append(record)
            return records
